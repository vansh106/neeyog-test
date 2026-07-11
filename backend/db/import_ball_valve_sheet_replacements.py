"""Replace all ball valve catalog sheets + supplier prices from replacement workbook.

Workbook: ``docs/Final_Products/replacement/Ball_Valve_Products (1).xlsx``

Usage (from ``backend/``):
  python -m db.import_ball_valve_sheet_replacements --dry-run
  python -m db.import_ball_valve_sheet_replacements
"""

from __future__ import annotations

import argparse
import asyncio
import logging
import sys
import uuid
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import delete, func, select

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from core.config import get_settings  # noqa: E402
from core.database import async_session_factory, init_db  # noqa: E402
from db.ball_valve_sheet_constants import (  # noqa: E402
    BALL_VALVE_CATALOG_KEYS,
    BALL_VALVE_REPLACEMENT_FILENAME,
    BALL_VALVE_REPLACEMENT_SHEETS,
    CASCO_SUPPLIER_NAME,
    UNISON_SUPPLIER_NAME,
)
from db.final_product_models import FINAL_PRODUCT_SHEET_MODELS  # noqa: E402
from db.import_final_products_catalog import (  # noqa: E402
    DOCS_FINAL,
    _build_header_map,
    _cell_float,
    _cell_str,
    _header_to_field,
)
from db.models import Supplier, SupplierProductPrice  # noqa: E402
from services.pricing_service import match_catalog_row  # noqa: E402

logger = logging.getLogger(__name__)

MODEL_BY_KEY: dict[str, type] = dict(FINAL_PRODUCT_SHEET_MODELS)
REPLACEMENT_DIR = DOCS_FINAL / "replacement"


def _price_column_index(header_row: tuple[Any, ...]) -> int | None:
    for i, cell in enumerate(header_row):
        if _header_to_field(cell) == "price":
            return i
    return None


def _apply_row_fields(
    kwargs: dict[str, Any],
    data_row: tuple[Any, ...],
    idx_map: dict[int, str],
) -> bool:
    empty = True
    for col_idx, field in idx_map.items():
        if col_idx >= len(data_row):
            continue
        raw = data_row[col_idx]
        if field == "sr_no":
            continue
        kwargs[field] = _cell_str(raw)
        if field != "sr_no" and kwargs.get(field) not in (None, ""):
            empty = False
    return not empty


def _row_spec_and_price(
    data_row: tuple[Any, ...],
    idx_map: dict[int, str],
    price_idx: int | None,
) -> tuple[dict[str, str], float | None]:
    spec: dict[str, str] = {}
    for col_idx, field in idx_map.items():
        if col_idx >= len(data_row):
            continue
        if field == "sr_no":
            continue
        s = _cell_str(data_row[col_idx])
        if s:
            spec[field] = s

    price: float | None = None
    if price_idx is not None and price_idx < len(data_row):
        price = _cell_float(data_row[price_idx])
    return spec, price


async def _find_supplier(session: Any, client_id: str, name_hint: str) -> Supplier | None:
    hint = name_hint.strip().lower()
    result = await session.execute(
        select(Supplier).where(
            Supplier.client_id == client_id,
            Supplier.is_active.is_(True),
            func.lower(Supplier.name) == hint,
        )
    )
    row = result.scalar_one_or_none()
    if row:
        return row
    result = await session.execute(
        select(Supplier).where(
            Supplier.client_id == client_id,
            Supplier.is_active.is_(True),
            func.lower(Supplier.name).contains(hint),
        )
    )
    return result.scalars().first()


async def _upsert_price(
    session: Any,
    supplier: Supplier,
    catalog_key: str,
    spec: dict[str, str],
    price_f: float,
    stats: dict[str, int],
) -> None:
    if price_f <= 0:
        stats["no_price"] += 1
        return
    rid, reason = await match_catalog_row(catalog_key, spec, session)
    if rid is None:
        stats["unmatched"] += 1
        logger.warning("Price unmatched %s (%s): %s", catalog_key, reason, spec)
        return

    result = await session.execute(
        select(SupplierProductPrice).where(
            SupplierProductPrice.supplier_id == supplier.id,
            SupplierProductPrice.catalog_table == catalog_key,
            SupplierProductPrice.catalog_row_id == uuid.UUID(rid),
        )
    )
    existing = result.scalar_one_or_none()
    if existing:
        existing.list_price_inr = float(price_f)
        stats["updated"] += 1
    else:
        session.add(
            SupplierProductPrice(
                supplier_id=supplier.id,
                catalog_table=catalog_key,
                catalog_row_id=uuid.UUID(rid),
                list_price_inr=float(price_f),
                discount_pct_override=None,
            )
        )
        stats["created"] += 1


async def _import_sheet(
    session: Any,
    *,
    wb_path: Path,
    sheet_name: str,
    catalog_key: str,
    source_file: str,
    client_id: str,
    supplier: Supplier,
    stats: dict[str, int],
) -> int:
    model = MODEL_BY_KEY.get(catalog_key)
    if model is None:
        raise ValueError(f"Unknown catalog_key: {catalog_key}")

    wb = load_workbook(wb_path, read_only=True, data_only=True)
    if sheet_name not in wb.sheetnames:
        wb.close()
        raise ValueError(f"Sheet {sheet_name!r} not in {wb_path.name}; have: {wb.sheetnames}")

    ws = wb[sheet_name]
    rows_iter = ws.iter_rows(values_only=True)
    header_row = next(rows_iter, ())
    idx_map = _build_header_map(header_row, model)
    price_idx = _price_column_index(header_row)
    if not idx_map:
        wb.close()
        raise ValueError(f"No mappable headers for sheet {sheet_name!r}")

    n_cat = 0
    for data_row in rows_iter:
        if not data_row or not any(c is not None and str(c).strip() != "" for c in data_row):
            continue
        kwargs: dict[str, Any] = {
            "row_id": uuid.uuid4(),
            "client_id": client_id,
            "source_file": source_file,
        }
        if not _apply_row_fields(kwargs, data_row, idx_map):
            continue
        n_cat += 1
        kwargs["sr_no"] = float(n_cat)
        session.add(model(**kwargs))

        spec, price_f = _row_spec_and_price(data_row, idx_map, price_idx)
        if price_f is None:
            stats["no_price"] += 1
        else:
            await _upsert_price(session, supplier, catalog_key, spec, price_f, stats)

    wb.close()
    await session.flush()
    logger.info(
        "Imported %d catalog rows -> %s (%s, supplier=%s)",
        n_cat,
        model.__tablename__,
        sheet_name,
        supplier.name,
    )
    return n_cat


async def _run(dry_run: bool) -> None:
    settings = get_settings()
    client_id = settings.ACTIVE_CLIENT
    wb_path = REPLACEMENT_DIR / BALL_VALVE_REPLACEMENT_FILENAME

    if not wb_path.is_file():
        raise FileNotFoundError(f"Missing ball valve workbook: {wb_path}")

    if dry_run:
        logger.info("Dry run — client_id=%s", client_id)
        wb = load_workbook(wb_path, read_only=True, data_only=True)
        for sheet_name, catalog_key, supplier_name in BALL_VALVE_REPLACEMENT_SHEETS:
            if sheet_name not in wb.sheetnames:
                logger.warning("  MISSING sheet %r for %s", sheet_name, catalog_key)
                continue
            ws = wb[sheet_name]
            n = sum(
                1
                for r in ws.iter_rows(min_row=2, values_only=True)
                if r and any(c is not None and str(c).strip() for c in r)
            )
            logger.info(
                "  %s -> %s rows=%d supplier=%s",
                sheet_name,
                catalog_key,
                n,
                supplier_name,
            )
        wb.close()
        logger.info("Will clear catalog keys: %s", sorted(BALL_VALVE_CATALOG_KEYS))
        return

    await init_db()
    async with async_session_factory() as session:
        casco = await _find_supplier(session, client_id, CASCO_SUPPLIER_NAME)
        unison = await _find_supplier(session, client_id, UNISON_SUPPLIER_NAME)
        if casco is None:
            raise RuntimeError(f"No active supplier matching {CASCO_SUPPLIER_NAME!r}")
        if unison is None:
            raise RuntimeError(f"No active supplier matching {UNISON_SUPPLIER_NAME!r}")
        logger.info("Casco supplier: %s (%s)", casco.name, casco.id)
        logger.info("Unison supplier: %s (%s)", unison.name, unison.id)

        suppliers_by_name = {
            CASCO_SUPPLIER_NAME.lower(): casco,
            UNISON_SUPPLIER_NAME.lower(): unison,
        }

        models = [MODEL_BY_KEY[k] for k in sorted(BALL_VALVE_CATALOG_KEYS) if k in MODEL_BY_KEY]
        for model in models:
            await session.execute(delete(model).where(model.client_id == client_id))
        await session.execute(
            delete(SupplierProductPrice).where(
                SupplierProductPrice.catalog_table.in_(sorted(BALL_VALVE_CATALOG_KEYS)),
            )
        )
        await session.flush()

        catalog_counts: dict[str, int] = {}
        stats = {"created": 0, "updated": 0, "unmatched": 0, "no_price": 0}

        for sheet_name, catalog_key, supplier_name in BALL_VALVE_REPLACEMENT_SHEETS:
            supplier = suppliers_by_name.get(supplier_name.lower())
            if supplier is None:
                raise RuntimeError(f"No supplier for {supplier_name!r}")
            n = await _import_sheet(
                session,
                wb_path=wb_path,
                sheet_name=sheet_name,
                catalog_key=catalog_key,
                source_file=BALL_VALVE_REPLACEMENT_FILENAME,
                client_id=client_id,
                supplier=supplier,
                stats=stats,
            )
            catalog_counts[catalog_key] = n

        await session.commit()
        logger.info("Catalog sheets: %s", catalog_counts)
        logger.info(
            "Supplier prices: created=%d updated=%d unmatched=%d no_price=%d",
            stats["created"],
            stats["updated"],
            stats["unmatched"],
            stats["no_price"],
        )


def main() -> None:
    logging.basicConfig(level=logging.INFO)
    p = argparse.ArgumentParser(
        description="Replace ball valve masters from Ball_Valve_Products (1).xlsx",
    )
    p.add_argument("--dry-run", action="store_true")
    args = p.parse_args()
    asyncio.run(_run(dry_run=args.dry_run))


if __name__ == "__main__":
    main()
