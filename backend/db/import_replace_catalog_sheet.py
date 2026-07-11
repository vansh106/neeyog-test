"""Replace one ``catalog_fp_*`` sheet from a standalone Excel file.

Clears existing rows (and supplier prices for that catalog_table), then imports rows.

Usage (from ``backend/``):
  python -m db.import_replace_catalog_sheet \\
    --catalog-key fp_ball_valve_casco_1_piece_multi_end \\
    --file ../docs/Final_Products/replacement/1-Piece\\ Ball\\ Valve\\ Casco\\ make.xlsx \\
    --sheet "Casco – 1-Piece Ball Valve (2)"
"""

from __future__ import annotations

import argparse
import asyncio
import logging
import sys
import uuid
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import delete

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from core.config import get_settings  # noqa: E402
from core.database import async_session_factory, init_db  # noqa: E402
from db.final_product_models import FINAL_PRODUCT_SHEET_MODELS  # noqa: E402
from db.import_final_products_catalog import (  # noqa: E402
    _build_header_map,
    _cell_float,
    _cell_str,
)
from db.models import SupplierProductPrice  # noqa: E402

logger = logging.getLogger(__name__)

MODEL_BY_KEY: dict[str, type] = dict(FINAL_PRODUCT_SHEET_MODELS)


def _import_rows_from_workbook(
    path: Path,
    sheet_name: str,
    model: type,
    client_id: str,
    source_file: str,
) -> list[Any]:
    wb = load_workbook(path, read_only=True, data_only=True)
    if sheet_name not in wb.sheetnames:
        wb.close()
        raise ValueError(f"Sheet {sheet_name!r} not in {path.name}; have: {wb.sheetnames}")
    ws = wb[sheet_name]
    rows_iter = ws.iter_rows(min_row=1, values_only=True)
    header_row = next(rows_iter, ())
    idx_map = _build_header_map(header_row, model)
    if not idx_map:
        wb.close()
        raise ValueError(f"No mappable headers in {path.name} / {sheet_name}")

    instances: list[Any] = []
    for data_row in rows_iter:
        kwargs: dict[str, Any] = {}
        kwargs["row_id"] = uuid.uuid4()
        kwargs["client_id"] = client_id
        kwargs["source_file"] = source_file
        empty = True
        for col_idx, field in idx_map.items():
            if col_idx >= len(data_row):
                continue
            raw = data_row[col_idx]
            if field == "sr_no":
                kwargs[field] = _cell_float(raw)
            else:
                kwargs[field] = _cell_str(raw)
            if kwargs.get(field) not in (None, ""):
                empty = False
        if empty:
            continue
        instances.append(model(**kwargs))
    wb.close()
    return instances


async def replace_catalog_sheet(
    catalog_key: str,
    xlsx_path: Path,
    sheet_name: str | None,
    *,
    dry_run: bool = False,
    source_file_label: str | None = None,
) -> int:
    model = MODEL_BY_KEY.get(catalog_key)
    if model is None:
        raise ValueError(f"Unknown catalog_key: {catalog_key}")

    if not xlsx_path.is_file():
        raise FileNotFoundError(xlsx_path)

    wb_probe = load_workbook(xlsx_path, read_only=True, data_only=True)
    resolved_sheet = sheet_name or wb_probe.sheetnames[0]
    wb_probe.close()

    settings = get_settings()
    client_id = settings.ACTIVE_CLIENT
    source_file = source_file_label or xlsx_path.name

    instances = _import_rows_from_workbook(
        xlsx_path, resolved_sheet, model, client_id, source_file
    )

    if dry_run:
        logger.info(
            "Dry run: would replace %s with %d rows from %s (%s)",
            catalog_key,
            len(instances),
            xlsx_path.name,
            resolved_sheet,
        )
        return len(instances)

    await init_db()
    async with async_session_factory() as session:
        del_cat = await session.execute(
            delete(model).where(model.client_id == client_id)
        )
        del_prices = await session.execute(
            delete(SupplierProductPrice).where(
                SupplierProductPrice.catalog_table == catalog_key,
            )
        )
        for row in instances:
            session.add(row)
        await session.commit()
        logger.info(
            "Replaced %s: removed %d catalog rows, %d supplier prices; inserted %d rows",
            catalog_key,
            del_cat.rowcount or 0,
            del_prices.rowcount or 0,
            len(instances),
        )
    return len(instances)


def main() -> None:
    logging.basicConfig(level=logging.INFO)
    p = argparse.ArgumentParser(description="Replace one masters catalog sheet from Excel")
    p.add_argument("--catalog-key", required=True)
    p.add_argument("--file", required=True, type=Path)
    p.add_argument("--sheet", default=None, help="Worksheet name (default: first sheet)")
    p.add_argument("--source-file-label", default=None, help="Stored in source_file column")
    p.add_argument("--dry-run", action="store_true")
    args = p.parse_args()

    path = args.file
    if not path.is_absolute():
        path = (Path.cwd() / path).resolve()

    asyncio.run(
        replace_catalog_sheet(
            args.catalog_key,
            path,
            args.sheet,
            dry_run=args.dry_run,
            source_file_label=args.source_file_label,
        )
    )


if __name__ == "__main__":
    main()
