"""Import Aluminium Foil products from Wholesale Price List workbook.

Parses worksheet ``Aluminium Foil`` — product specs go to ``catalog_fp_aluminium_foil``;
list prices are stored per supplier in ``supplier_product_prices`` (suppliers are not
embedded in catalog rows).

Usage (from ``backend/``):
  python -m db.import_aluminium_foil_catalog_and_prices --dry-run
  python -m db.import_aluminium_foil_catalog_and_prices
"""

from __future__ import annotations

import argparse
import asyncio
import logging
import re
import sys
import uuid
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

logger = logging.getLogger(__name__)

WORKBOOK = "Wholesale Price List Date 16.06.2026.xlsx"
WORKSHEET = "Aluminium Foil"
CATALOG_KEY = "fp_aluminium_foil"
DOCS_FINAL = Path(__file__).resolve().parents[2] / "docs" / "Final_Products"

CAT_FOIL_WRAP = "Foil Wrap"
CAT_FOIL_BOX = "Foil Box"
CAT_FOIL_CONTAINER = "Foil Container"
CAT_PREMIUM_CONTAINER = "Premium Foil Container"
CAT_FOIL_PAPER_LIDS = "Foil Paper Lids"
CAT_EXCLUSIVE_CONTAINER = "Exclusive Foil Container"
CAT_PET_LID = "Pet Lid"

ALL_CATEGORIES = (
    CAT_FOIL_WRAP,
    CAT_FOIL_BOX,
    CAT_FOIL_CONTAINER,
    CAT_PREMIUM_CONTAINER,
    CAT_FOIL_PAPER_LIDS,
    CAT_EXCLUSIVE_CONTAINER,
    CAT_PET_LID,
)

SUPPLIER_NAMES = (
    "Hindalco Superwrap",
    "FW Foil",
    "Minda",
    "PNS",
    "WrappFresh",
    "Alu Fresh",
    "Pet Lid",
    "Foil Paper Lids",
    "Exclusive Foil",
)


@dataclass
class ParsedProduct:
    sr_no: int | None
    category: str
    supplier_name: str
    product_name: str
    size_value: str | None
    size_unit: str | None
    dimensions: str | None
    weight_grade: str | None
    product_code: str | None
    hsn_code: str | None
    std_pack: str | None
    pack_unit: str | None
    rate_incl_gst: float | None
    canonical_sku: str
    raw_description: str


def _cell_str(v: Any) -> str | None:
    if v is None:
        return None
    s = str(v).strip()
    return s or None


def _cell_float(v: Any) -> float | None:
    if v is None or v == "" or v == "-":
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def _is_section_header(row: tuple[Any, ...]) -> bool:
    sr, code, _hsn, desc = (row + (None,) * 4)[:4]
    if not isinstance(desc, str):
        return False
    d = desc.strip()
    if not d:
        return False
    if isinstance(sr, (int, float)):
        return False
    dl = d.lower()
    if "hsn code" in dl or d.endswith(")"):
        return True
    if "hindalco" in dl or "superwrap" in dl:
        return True
    if "pet lid" in dl and "foil cont" in dl:
        return True
    if "foil paper lids" in dl:
        return True
    return False


def _section_category_supplier(desc: str) -> tuple[str, str] | None:
    d = desc.strip()
    dl = d.lower()
    if "hindalco" in dl or "superwrap" in dl:
        return CAT_FOIL_WRAP, "Hindalco Superwrap"
    if "minda foil container premium" in dl:
        return CAT_PREMIUM_CONTAINER, "Minda"
    if "minda foil container" in dl:
        return CAT_FOIL_CONTAINER, "Minda"
    if "minda foil" in dl:
        return CAT_FOIL_WRAP, "Minda"
    if "pns foil" in dl:
        return CAT_FOIL_WRAP, "PNS"
    if "wrappfresh" in dl or "wrapp fresh" in dl:
        return CAT_FOIL_WRAP, "WrappFresh"
    if "alufresh" in dl or "alu fresh" in dl:
        return CAT_FOIL_CONTAINER, "Alu Fresh"
    if "pet lid for foil" in dl:
        return CAT_PET_LID, "Pet Lid"
    if "foil paper lids" in dl:
        return CAT_FOIL_PAPER_LIDS, "Foil Paper Lids"
    if "exclusive products" in dl:
        return CAT_EXCLUSIVE_CONTAINER, "Exclusive Foil"
    return None


def _infer_supplier_from_description(desc: str) -> str:
    dl = desc.lower()
    if "(mi)" in dl or "minda" in dl:
        return "Minda"
    if "(sw)" in dl or "superwrap" in dl:
        return "Hindalco Superwrap"
    if "(fw)" in dl:
        return "FW Foil"
    if "pns" in dl:
        return "PNS"
    if "wrapp" in dl:
        return "WrappFresh"
    if "alu fresh" in dl or "alufresh" in dl:
        return "Alu Fresh"
    return "Exclusive Foil"


def _extract_size(desc: str, category: str) -> tuple[str | None, str | None]:
    d = desc
    for pat, unit in (
        (r"(\d+(?:\.\d+)?)\s*ML\b", "ml"),
        (r"(\d+(?:\.\d+)?)\s*MTR\b", "mtr"),
        (r"(\d+(?:\.\d+)?)\s*KG\b", "kg"),
        (r"(\d+(?:\.\d+)?)\s*LTR\b", "ltr"),
        (r"(\d+(?:\.\d+)?)\s*No\b", "no"),
        (r"(\d+(?:\.\d+)?)\s*\"\s*Round", "inch"),
        (r"(\d+(?:\.\d+)?)\s*INCH\b", "inch"),
        (r"(\d+(?:\.\d+)?)\s*CP\b", "cp"),
    ):
        m = re.search(pat, d, re.I)
        if m:
            val = m.group(1)
            if unit == "ml":
                return f"{val}ML", unit
            if unit == "mtr":
                return f"{val}MTR", unit
            if unit == "kg":
                return f"{val}KG", unit
            if unit == "inch":
                return f'{val}"', unit
            if unit == "cp":
                return f"{val}CP", unit
            return val, unit
  # AFC codes with embedded ml
    m = re.search(r"AFC\s*\d+\((\d+)ml\)", d, re.I)
    if m:
        return f"{m.group(1)}ML", "ml"
    m = re.search(r"(\d+)ml\b", d, re.I)
    if m and category != CAT_FOIL_WRAP:
        return f"{m.group(1)}ML", "ml"
    return None, None


def _weight_grade(desc: str) -> str | None:
    dl = desc.lower()
    if "light" in dl:
        return "Light"
    if "heavy" in dl:
        return "Heavy"
    return None


def _clean_product_name(desc: str, category: str) -> str:
    s = desc.strip()
    # Strip embedded supplier / brand tokens from product title.
    patterns = [
        r"\bMoments\b",
        r"\b/?\s*Alu\s*Fresh\b",
        r"\bAlufresh\b",
        r"\bMinda\b",
        r"\bPNS\b",
        r"\bWrapp\s*Fresh\b",
        r"\bWrappFresh\b",
        r"\bSuperwrap\b",
        r"\bHindalco\b",
        r"\(MI\)",
        r"\(SW\)",
        r"\(FW\)",
        r"\(NG\)",
        r"\(HI\)",
        r"\(NW\)",
        r"\(GW\)",
        r"\s*\*\s*\d+\s*pcs",
        r"\s+",
    ]
    for pat in patterns:
        s = re.sub(pat, " ", s, flags=re.I)
    s = re.sub(r"\s+", " ", s).strip(" /-")
    if category == CAT_FOIL_WRAP and "wrap" not in s.lower():
        s = f"{s} Foil Wrap".strip()
    if category == CAT_FOIL_CONTAINER and "cont" not in s.lower():
        s = f"{s} Foil Container".strip()
    return s or desc.strip()


def _canonical_sku(
    category: str,
    product_name: str,
    size_value: str | None,
    dimensions: str | None,
    weight_grade: str | None,
) -> str:
    parts = [
        category.lower(),
        (size_value or "").lower(),
        (dimensions or "").lower().replace(" ", ""),
        (weight_grade or "").lower(),
        re.sub(r"[^a-z0-9]+", "-", product_name.lower()).strip("-"),
    ]
    return "|".join(p for p in parts if p)


def _parse_dimensions(code: Any) -> str | None:
    if not isinstance(code, str):
        return None
    c = code.strip()
    if re.match(r"^\d+\s*[xX]\s*\d+", c):
        return c.upper().replace(" ", "")
    return None


def _parse_workbook(path: Path) -> list[ParsedProduct]:
    wb = load_workbook(path, read_only=True, data_only=True)
    if WORKSHEET not in wb.sheetnames:
        raise ValueError(f"Worksheet {WORKSHEET!r} not found")
    ws = wb[WORKSHEET]

    current_category = CAT_FOIL_WRAP
    current_supplier = "Hindalco Superwrap"
    out: list[ParsedProduct] = []

    for row in ws.iter_rows(values_only=True):
        if len(row) < 4:
            continue
        sr, code, hsn, desc = row[0], row[1], row[2], row[3]
        if isinstance(desc, str) and _is_section_header(row):
            parsed = _section_category_supplier(desc)
            if parsed:
                current_category, current_supplier = parsed
            continue

        if not isinstance(desc, str) or not desc.strip():
            continue
        if isinstance(code, str) and re.match(r"^FR\d+FW$", code.strip(), re.I):
            current_category = CAT_FOIL_WRAP
            current_supplier = "FW Foil"

        if not isinstance(sr, (int, float)):
            continue

        description = desc.strip()
        dl = description.lower()
        category = current_category
        if "box" in dl and ("wrap" in dl or "foil" in dl):
            category = CAT_FOIL_BOX

        supplier = current_supplier
        if category == CAT_EXCLUSIVE_CONTAINER:
            supplier = _infer_supplier_from_description(description)

        size_value, size_unit = _extract_size(description, category)
        dimensions = _parse_dimensions(code)
        weight = _weight_grade(description)
        cleaned = _clean_product_name(description, category)
        sku = _canonical_sku(category, cleaned, size_value, dimensions, weight)

        rate_inc = _cell_float(row[5] if len(row) > 5 else None)
        std_pack = _cell_str(row[7] if len(row) > 7 else None)
        pack_unit = _cell_str(row[6] if len(row) > 6 else None)

        out.append(
            ParsedProduct(
                sr_no=int(sr) if isinstance(sr, (int, float)) and sr == int(sr) else None,
                category=category,
                supplier_name=supplier,
                product_name=cleaned,
                size_value=size_value,
                size_unit=size_unit,
                dimensions=dimensions,
                weight_grade=weight,
                product_code=_cell_str(code) if not dimensions else _cell_str(code),
                hsn_code=_cell_str(hsn),
                std_pack=std_pack,
                pack_unit=pack_unit,
                rate_incl_gst=rate_inc,
                canonical_sku=sku,
                raw_description=description,
            )
        )

    wb.close()
    return out


async def _get_or_create_supplier(session: Any, client_id: str, name: str) -> Any:
    from sqlalchemy import func, select

    from db.models import Supplier

    result = await session.execute(
        select(Supplier).where(
            Supplier.client_id == client_id,
            func.lower(Supplier.name) == name.strip().lower(),
        )
    )
    row = result.scalar_one_or_none()
    if row:
        return row
    supplier = Supplier(
        client_id=client_id,
        name=name.strip(),
        primary_category_key=CATALOG_KEY,
        is_active=True,
    )
    session.add(supplier)
    await session.flush()
    logger.info("Created supplier %s", name)
    return supplier


async def _run(dry_run: bool) -> None:
    path = DOCS_FINAL / WORKBOOK
    if not path.is_file():
        raise FileNotFoundError(path)

    products = _parse_workbook(path)
    logger.info("Parsed %d product lines from %s", len(products), WORKSHEET)

    by_cat: dict[str, int] = {}
    for p in products:
        by_cat[p.category] = by_cat.get(p.category, 0) + 1
    for cat in ALL_CATEGORIES:
        logger.info("  %s: %d", cat, by_cat.get(cat, 0))

    if dry_run:
        for p in products[:8]:
            logger.info(
                "  [%s / %s] %s size=%s %s price=%s",
                p.category,
                p.supplier_name,
                p.product_name,
                p.size_value,
                p.weight_grade or "",
                p.rate_incl_gst,
            )
        return

    from sqlalchemy import delete, func, select

    from core.config import get_settings
    from core.database import async_session_factory, init_db
    from db.final_product_models import CatalogFpAluminiumFoilRow
    from db.models import Supplier, SupplierProductPrice

    settings = get_settings()
    client_id = settings.ACTIVE_CLIENT
    await init_db()
    async with async_session_factory() as session:
        await session.execute(delete(CatalogFpAluminiumFoilRow).where(CatalogFpAluminiumFoilRow.client_id == client_id))
        await session.execute(
            delete(SupplierProductPrice).where(SupplierProductPrice.catalog_table == CATALOG_KEY)
        )
        await session.flush()

        sku_to_row_id: dict[str, uuid.UUID] = {}
        price_stats = {"created": 0, "updated": 0, "skipped": 0}

        for p in products:
            row_id = sku_to_row_id.get(p.canonical_sku)
            if row_id is None:
                row_id = uuid.uuid4()
                session.add(
                    CatalogFpAluminiumFoilRow(
                        row_id=row_id,
                        client_id=client_id,
                        sr_no=float(p.sr_no) if p.sr_no is not None else None,
                        variant_type=p.category,
                        product_name=p.product_name,
                        size_value=p.size_value,
                        size_unit=p.size_unit,
                        dimensions=p.dimensions,
                        weight_grade=p.weight_grade,
                        product_code=p.product_code,
                        hsn_code=p.hsn_code,
                        std_pack=p.std_pack,
                        pack_unit=p.pack_unit,
                        canonical_sku=p.canonical_sku,
                        source_file=WORKBOOK,
                    )
                )
                sku_to_row_id[p.canonical_sku] = row_id

            if p.rate_incl_gst is None:
                price_stats["skipped"] += 1
                continue

            supplier = await _get_or_create_supplier(session, client_id, p.supplier_name)
            result = await session.execute(
                select(SupplierProductPrice).where(
                    SupplierProductPrice.supplier_id == supplier.id,
                    SupplierProductPrice.catalog_table == CATALOG_KEY,
                    SupplierProductPrice.catalog_row_id == row_id,
                )
            )
            existing = result.scalar_one_or_none()
            if existing:
                existing.list_price_inr = float(p.rate_incl_gst)
                price_stats["updated"] += 1
            else:
                session.add(
                    SupplierProductPrice(
                        supplier_id=supplier.id,
                        catalog_table=CATALOG_KEY,
                        catalog_row_id=row_id,
                        list_price_inr=float(p.rate_incl_gst),
                    )
                )
                price_stats["created"] += 1

        await session.commit()

        cat_n = (
            await session.execute(
                select(func.count()).select_from(CatalogFpAluminiumFoilRow).where(
                    CatalogFpAluminiumFoilRow.client_id == client_id
                )
            )
        ).scalar_one()
        price_n = (
            await session.execute(
                select(func.count())
                .select_from(SupplierProductPrice)
                .where(SupplierProductPrice.catalog_table == CATALOG_KEY)
            )
        ).scalar_one()
        logger.info("Imported catalog rows=%s supplier_prices=%s", cat_n, price_n)
        logger.info("Price stats: %s", price_stats)


def main() -> None:
    logging.basicConfig(level=logging.INFO)
    p = argparse.ArgumentParser()
    p.add_argument("--dry-run", action="store_true")
    args = p.parse_args()
    asyncio.run(_run(dry_run=args.dry_run))


if __name__ == "__main__":
    main()
