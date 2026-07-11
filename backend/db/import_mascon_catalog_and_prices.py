"""Re-import Mascon catalog from ``Mascon_Valve_Products (1) - F (1).xlsx`` and load list prices for supplier Mascon.

Clears only ``catalog_fp_mascon_*`` rows for ACTIVE_CLIENT, re-imports all 13 worksheets,
removes prior Mascon supplier prices for those categories, then upserts prices from the Price column.

Usage (from ``backend/``):
  python -m db.import_mascon_catalog_and_prices --dry-run
  python -m db.import_mascon_catalog_and_prices
  python -m db.import_mascon_catalog_and_prices --supplier-name "Mascon"
"""

from __future__ import annotations

import argparse
import asyncio
import logging
import sys
import uuid
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import delete, func, select

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from core.config import get_settings  # noqa: E402
from core.database import async_session_factory, init_db  # noqa: E402
from db import final_product_models as fpm  # noqa: E402
from db.import_final_products_catalog import (  # noqa: E402
    DOCS_FINAL,
    _build_header_map,
    _cell_float,
    _cell_str,
    _header_to_field,
)
from db.models import Supplier, SupplierProductPrice  # noqa: E402
from services.pricing_service import match_catalog_row  # noqa: E402

logger = logging.getLogger(__name__)

MASCON_FILENAME = "Mascon_Valve_Products (1) - F (1).xlsx"

# (worksheet name, API catalog_table key, ORM model)
MASCON_SHEETS: list[tuple[str, str, type]] = [
    ("Manual – TC End", "fp_mascon_manual_tc_end", fpm.CatalogFpMasconManualTcEndRow),
    ("Manual – Butt Weld", "fp_mascon_manual_butt_weld", fpm.CatalogFpMasconManualButtWeldRow),
    ("Pneumatic – TC End", "fp_mascon_pneumatic_tc_end", fpm.CatalogFpMasconPneumaticTcEndRow),
    ("Pneumatic – Butt Weld", "fp_mascon_pneumatic_butt_weld", fpm.CatalogFpMasconPneumaticButtWeldRow),
    ("ZDV-M – L Type", "fp_mascon_zdvm_l_type", fpm.CatalogFpMasconZdvmLTypeRow),
    ("ZDV-M – J Type", "fp_mascon_zdvm_j_type", fpm.CatalogFpMasconZdvmJTypeRow),
    ("ZDV-P – L Type", "fp_mascon_zdvp_l_type", fpm.CatalogFpMasconZdvpLTypeRow),
    ("ZDV-P – J Type", "fp_mascon_zdvp_j_type", fpm.CatalogFpMasconZdvpJTypeRow),
    ("PRV", "fp_mascon_prv", fpm.CatalogFpMasconPrvRow),
    ("Angle – Screwed & Flanged", "fp_mascon_angle_sc_flanged", fpm.CatalogFpMasconAngleScFlangedRow),
    ("Angle – Butt Weld", "fp_mascon_angle_butt_weld", fpm.CatalogFpMasconAngleButtWeldRow),
    ("Angle – TC End", "fp_mascon_angle_tc_end", fpm.CatalogFpMasconAngleTcEndRow),
    ("Spare Diaphragm", "fp_mascon_spare_diaphragm", fpm.CatalogFpMasconSpareDiaphragmRow),
]

MASCON_CATALOG_KEYS = [k for _, k, _ in MASCON_SHEETS]


def _price_column_index(header_row: tuple[Any, ...]) -> int | None:
    for i, cell in enumerate(header_row):
        if _header_to_field(cell) == "price":
            return i
    return None


def _row_spec_and_price(
    data_row: tuple[Any, ...],
    idx_map: dict[int, str],
    price_idx: int | None,
) -> tuple[dict[str, str], float | None]:
  spec: dict[str, str] = {}
  for col_idx, field in idx_map.items():
    if col_idx >= len(data_row):
      continue
    raw = data_row[col_idx]
    if field == "sr_no":
      continue
    s = _cell_str(raw)
    if s:
      spec[field] = s
  price: float | None = None
  if price_idx is not None and price_idx < len(data_row):
    price = _cell_float(data_row[price_idx])
  return spec, price


async def _find_supplier(session: Any, client_id: str, name_hint: str) -> Supplier | None:
    hint = name_hint.strip().lower()
    result = await session.execute(
        select(Supplier).where(
            Supplier.client_id == client_id,
            Supplier.is_active.is_(True),
            func.lower(Supplier.name) == hint,
        )
    )
    row = result.scalar_one_or_none()
    if row:
        return row
    result = await session.execute(
        select(Supplier).where(
            Supplier.client_id == client_id,
            Supplier.is_active.is_(True),
            func.lower(Supplier.name).contains(hint),
        )
    )
    return result.scalars().first()


async def _run(dry_run: bool, supplier_name: str) -> None:
    settings = get_settings()
    client_id = settings.ACTIVE_CLIENT
    path = DOCS_FINAL / MASCON_FILENAME
    if not path.is_file():
        raise FileNotFoundError(f"Missing workbook: {path}")

    if dry_run:
        wb = load_workbook(path, read_only=True, data_only=True)
        logger.info("Dry run — would re-import Mascon catalog + prices for client_id=%s", client_id)
        for sheet_name, catalog_key, model in MASCON_SHEETS:
            if sheet_name not in wb.sheetnames:
                logger.warning("  missing sheet %r", sheet_name)
                continue
            ws = wb[sheet_name]
            header = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
            price_ix = _price_column_index(header)
            n = sum(1 for _ in ws.iter_rows(min_row=2, values_only=True) if any(c is not None for c in _))
            logger.info("  %s -> %s (%s rows, price_col=%s)", sheet_name, catalog_key, n, price_ix)
        wb.close()
        return

    await init_db()
    async with async_session_factory() as session:
        supplier = await _find_supplier(session, client_id, supplier_name)
        if supplier is None:
            raise RuntimeError(
                f"No active supplier matching {supplier_name!r} for client {client_id}. "
                "Create the supplier in Masters first."
            )
        logger.info("Using supplier %s (%s)", supplier.name, supplier.id)

        mascon_models = [m for _, _, m in MASCON_SHEETS]
        for model in mascon_models:
            await session.execute(delete(model).where(model.client_id == client_id))
        await session.execute(
            delete(SupplierProductPrice).where(
                SupplierProductPrice.supplier_id == supplier.id,
                SupplierProductPrice.catalog_table.in_(MASCON_CATALOG_KEYS),
            )
        )
        await session.flush()

        wb = load_workbook(path, read_only=True, data_only=True)
        catalog_counts: dict[str, int] = {}
        price_stats = {"created": 0, "updated": 0, "unmatched": 0, "no_price": 0}

        for sheet_name, catalog_key, model in MASCON_SHEETS:
            if sheet_name not in wb.sheetnames:
                logger.warning("Skip missing sheet %r", sheet_name)
                continue
            ws = wb[sheet_name]
            rows_iter = ws.iter_rows(values_only=True)
            header_row = next(rows_iter, ())
            idx_map = _build_header_map(header_row, model)
            price_idx = _price_column_index(header_row)
            if not idx_map:
                logger.warning("No mappable headers for %s", sheet_name)
                continue
            if price_idx is None:
                logger.warning("No Price column on sheet %s", sheet_name)

            sheet_rows: list[tuple[Any, ...]] = []
            n_cat = 0
            for data_row in rows_iter:
                if not data_row or not any(c is not None and str(c).strip() != "" for c in data_row):
                    continue
                kwargs: dict[str, Any] = {
                    "row_id": uuid.uuid4(),
                    "client_id": client_id,
                    "source_file": MASCON_FILENAME,
                }
                empty = True
                for col_idx, field in idx_map.items():
                    if col_idx >= len(data_row):
                        continue
                    raw = data_row[col_idx]
                    if field == "sr_no":
                        kwargs[field] = _cell_float(raw)
                    else:
                        kwargs[field] = _cell_str(raw)
                    if field != "sr_no" and kwargs.get(field) not in (None, ""):
                        empty = False
                if empty:
                    continue
                session.add(model(**kwargs))
                sheet_rows.append(data_row)
                n_cat += 1

            await session.flush()

            for data_row in sheet_rows:
                spec, price_f = _row_spec_and_price(data_row, idx_map, price_idx)
                if price_f is None:
                    price_stats["no_price"] += 1
                    continue
                rid, reason = await match_catalog_row(catalog_key, spec, session)
                if rid is None:
                    price_stats["unmatched"] += 1
                    logger.warning("Price unmatched %s (%s): %s", catalog_key, reason, spec)
                    continue

                result = await session.execute(
                    select(SupplierProductPrice).where(
                        SupplierProductPrice.supplier_id == supplier.id,
                        SupplierProductPrice.catalog_table == catalog_key,
                        SupplierProductPrice.catalog_row_id == uuid.UUID(rid),
                    )
                )
                existing = result.scalar_one_or_none()
                if existing:
                    existing.list_price_inr = float(price_f)
                    price_stats["updated"] += 1
                else:
                    session.add(
                        SupplierProductPrice(
                            supplier_id=supplier.id,
                            catalog_table=catalog_key,
                            catalog_row_id=uuid.UUID(rid),
                            list_price_inr=float(price_f),
                            discount_pct_override=None,
                        )
                    )
                    price_stats["created"] += 1

            catalog_counts[catalog_key] = n_cat
            logger.info("Imported %d catalog rows -> %s (%s)", n_cat, model.__tablename__, sheet_name)

        wb.close()
        await session.commit()

        logger.info("Catalog sheets: %s", catalog_counts)
        logger.info(
            "Mascon prices: created=%d updated=%d unmatched=%d no_price=%d",
            price_stats["created"],
            price_stats["updated"],
            price_stats["unmatched"],
            price_stats["no_price"],
        )

        # Verify counts in DB
        for catalog_key, model in [(k, m) for _, k, m in MASCON_SHEETS]:
            cat_n = (
                await session.execute(
                    select(func.count()).select_from(model).where(model.client_id == client_id)
                )
            ).scalar_one()
            price_n = (
                await session.execute(
                    select(func.count())
                    .select_from(SupplierProductPrice)
                    .where(
                        SupplierProductPrice.supplier_id == supplier.id,
                        SupplierProductPrice.catalog_table == catalog_key,
                    )
                )
            ).scalar_one()
            logger.info("Verify %s: catalog=%s supplier_prices=%s", catalog_key, cat_n, price_n)


def main() -> None:
    logging.basicConfig(level=logging.INFO)
    p = argparse.ArgumentParser()
    p.add_argument("--dry-run", action="store_true")
    p.add_argument("--supplier-name", default="Mascon", help="Supplier name match (default: Mascon)")
    args = p.parse_args()
    asyncio.run(_run(dry_run=args.dry_run, supplier_name=args.supplier_name))


if __name__ == "__main__":
    main()
