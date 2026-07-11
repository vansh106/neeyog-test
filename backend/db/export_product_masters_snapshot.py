"""Export all Product Masters sheets from the database into a folder tree matching the UI.

Creates a timestamped directory under ``docs/``:
  ``docs/product_masters_snapshot_YYYY-MM-DD_HH-MM-SS/``

Each masters leaf (same paths as the frontend sidebar) becomes one ``.xlsx`` file.

Usage (from ``backend/``):
  python -m db.export_product_masters_snapshot

After changing ``frontend/lib/masterSidebarNav.ts``, refresh the leaf manifest:
  cd frontend && npx tsx scripts/export-master-nav-leaves.ts
"""

from __future__ import annotations

import argparse
import asyncio
import json
import logging
import re
import sys
from datetime import datetime
from pathlib import Path
from typing import Any

import sqlalchemy as sa
from openpyxl import Workbook
from openpyxl.styles import Font
from sqlalchemy import func, select

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from core.config import get_settings  # noqa: E402
from core.database import async_session_factory, init_db  # noqa: E402
from db.butterfly_sheet_constants import BUTTERFLY_NAV_SOURCE_FILE  # noqa: E402
from services.masters_service import (  # noqa: E402
    MASTERS_HIDDEN_SHEET_KEYS,
    MASTERS_SHEET_DISPLAY_COLUMNS,
    SHEET_MODEL_BY_KEY,
    _jsonable,
)

logger = logging.getLogger(__name__)

DOCS_DIR = Path(__file__).resolve().parents[2] / "docs"
NAV_LEAVES_JSON = Path(__file__).resolve().parent / "data" / "master_sidebar_nav_leaves.json"
LATEST_POINTER = DOCS_DIR / "product_masters_snapshot_LATEST.txt"

_SKIP_EXPORT_COLUMNS = frozenset({"row_id", "client_id", "created_at", "updated_at"})
_INVALID_PATH = re.compile(r'[<>:"/\\|?*\x00-\x1f]')


def _sanitize_segment(name: str) -> str:
    s = name.strip().replace("/", " - ")
    s = _INVALID_PATH.sub("_", s)
    return s.rstrip(". ") or "untitled"


def _column_header(field: str) -> str:
    return field.replace("_", " ").title()


def _export_columns(sheet_key: str, all_columns: list[str]) -> list[str]:
    preferred = MASTERS_SHEET_DISPLAY_COLUMNS.get(sheet_key)
    if preferred:
        return [c for c in preferred if c in all_columns]
    return [c for c in all_columns if c not in _SKIP_EXPORT_COLUMNS]


def _build_filters(leaf: dict[str, Any]) -> list:
    sheet = leaf["key"]
    model = SHEET_MODEL_BY_KEY.get(sheet)
    if model is None:
        raise ValueError(f"Unknown sheet key: {sheet}")

    filters: list = []
    vt = (leaf.get("variantType") or "").strip()
    if vt and hasattr(model, "variant_type"):
        filters.append(model.variant_type == vt)
    vc = (leaf.get("variantContains") or "").strip()
    if vc and hasattr(model, "variant_type"):
        filters.append(model.variant_type.ilike(f"%{vc}%"))
    vex = (leaf.get("variantExcludeContains") or "").strip()
    if vex and hasattr(model, "variant_type"):
        filters.append(~model.variant_type.ilike(f"%{vex}%"))
    vany = leaf.get("variantContainsAny")
    if vany and hasattr(model, "variant_type"):
        parts = [p.strip() for p in vany if p and str(p).strip()]
        if parts:
            filters.append(sa.or_(*[model.variant_type.ilike(f"%{p}%") for p in parts]))
    mnp = (leaf.get("modelNamePrefix") or "").strip()
    if mnp and hasattr(model, "model_name"):
        filters.append(model.model_name.ilike(f"{mnp}%"))
    nav = (leaf.get("navSlug") or "").strip()
    if sheet == "butterfly_valve" and nav and hasattr(model, "source_file"):
        source_label = BUTTERFLY_NAV_SOURCE_FILE.get(nav)
        if source_label:
            filters.append(model.source_file == source_label)
    return filters


async def _fetch_rows(session: Any, leaf: dict[str, Any], client_id: str) -> tuple[list[str], list[dict[str, Any]]]:
    sheet = leaf["key"]
    if sheet in MASTERS_HIDDEN_SHEET_KEYS:
        return [], []

    model = SHEET_MODEL_BY_KEY[sheet]
    filters = [model.client_id == client_id, *_build_filters(leaf)]

    order_parts = []
    if hasattr(model, "sr_no"):
        order_parts.append(model.sr_no.asc().nulls_last())
    if hasattr(model, "variant_type"):
        order_parts.append(model.variant_type.asc().nulls_last())
    if hasattr(model, "valve_size"):
        order_parts.append(model.valve_size.asc().nulls_last())
    if hasattr(model, "model_name"):
        order_parts.append(model.model_name.asc().nulls_last())
    order_parts.append(model.created_at.desc())

    rows = (await session.execute(select(model).where(*filters).order_by(*order_parts))).scalars().all()

    all_columns = list(model.__table__.columns.keys())
    columns = _export_columns(sheet, all_columns)
    items = [{c: _jsonable(getattr(r, c)) for c in columns} for r in rows]
    return columns, items


def _write_xlsx(path: Path, columns: list[str], items: list[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Products"
    header_font = Font(bold=True)
    for col_idx, field in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=col_idx, value=_column_header(field))
        cell.font = header_font
    for row_idx, item in enumerate(items, start=2):
        for col_idx, field in enumerate(columns, start=1):
            ws.cell(row=row_idx, column=col_idx, value=item.get(field))
    wb.save(path)


def _load_nav_leaves() -> list[dict[str, Any]]:
    if not NAV_LEAVES_JSON.is_file():
        raise FileNotFoundError(
            f"Missing {NAV_LEAVES_JSON}. Run: cd frontend && npx tsx scripts/export-master-nav-leaves.ts"
        )
    return json.loads(NAV_LEAVES_JSON.read_text(encoding="utf-8"))


async def _run(output_parent: Path | None = None) -> Path:
    leaves = _load_nav_leaves()
    stamp = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
    root = (output_parent or DOCS_DIR) / f"product_masters_snapshot_{stamp}"
    root.mkdir(parents=True, exist_ok=True)

    settings = get_settings()
    client_id = settings.ACTIVE_CLIENT

    await init_db()
    summary_lines = [
        f"Product Masters snapshot",
        f"Generated: {datetime.now().isoformat(timespec='seconds')}",
        f"Client: {client_id}",
        f"Source: database (ACTIVE_CLIENT)",
        f"Navigation: frontend/lib/masterSidebarNav.ts",
        "",
        "Sheet exports:",
    ]
    total_rows = 0

    async with async_session_factory() as session:
        for leaf in leaves:
            rel_dir = Path(*[_sanitize_segment(p) for p in leaf["path"]])
            filename = _sanitize_segment(leaf["label"]) + ".xlsx"
            out_path = root / rel_dir / filename

            columns, items = await _fetch_rows(session, leaf, client_id)
            _write_xlsx(out_path, columns, items)
            n = len(items)
            total_rows += n
            rel = out_path.relative_to(root)
            summary_lines.append(f"  {rel}  ({n} rows)")
            logger.info("Exported %s (%d rows)", rel, n)

    summary_lines.extend(["", f"Total rows: {total_rows}", f"Total sheets: {len(leaves)}"])
    readme = root / "README.txt"
    readme.write_text("\n".join(summary_lines) + "\n", encoding="utf-8")

    LATEST_POINTER.write_text(f"{root.name}\n{datetime.now().isoformat(timespec='seconds')}\n", encoding="utf-8")
    logger.info("Snapshot written to %s", root)
    logger.info("Latest pointer: %s", LATEST_POINTER)
    return root


def main() -> None:
    logging.basicConfig(level=logging.INFO)
    p = argparse.ArgumentParser(description="Export Product Masters from DB to docs/ folder tree")
    p.add_argument(
        "--output-parent",
        type=Path,
        default=None,
        help="Parent directory (default: repo docs/)",
    )
    args = p.parse_args()
    out_parent = args.output_parent
    if out_parent and not out_parent.is_absolute():
        out_parent = (Path.cwd() / out_parent).resolve()
    asyncio.run(_run(out_parent))


if __name__ == "__main__":
    main()
