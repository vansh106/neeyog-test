"""Backfill ``packing`` on existing strainer catalog rows from the replacement workbook.

Updates only the ``packing`` column; all other row data and supplier prices are unchanged.

Usage (from ``backend/``):
  python -m db.import_strainer_packing --dry-run
  python -m db.import_strainer_packing
"""

from __future__ import annotations

import argparse
import asyncio
import logging
import sys
import uuid
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import select

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from core.config import get_settings  # noqa: E402
from core.database import async_session_factory, init_db  # noqa: E402
from db.final_product_models import (  # noqa: E402
    CatalogFpStrainerY150Row,
    CatalogFpStrainerY300Row,
)
from db.import_final_products_catalog import (  # noqa: E402
    DOCS_FINAL,
    _build_header_map,
    _cell_str,
    _header_to_field,
)
from db.import_needle_sight_strainer_sheet_replacements import (  # noqa: E402
    _pressure_column_index,
    _strainer_catalog_key,
)
from db.needle_sight_strainer_sheet_constants import (  # noqa: E402
    STRAINER_REPLACEMENT_FILENAME,
    STRAINER_SHEET_NAME,
    STRAINER_Y150_KEY,
    STRAINER_Y300_KEY,
)
from services.pricing_service import match_catalog_row  # noqa: E402

logger = logging.getLogger(__name__)

REPLACEMENT_DIR = DOCS_FINAL / "replacement"
MODEL_BY_KEY = {
    STRAINER_Y150_KEY: CatalogFpStrainerY150Row,
    STRAINER_Y300_KEY: CatalogFpStrainerY300Row,
}


def _packing_column_index(header_row: tuple[Any, ...]) -> int | None:
    for i, cell in enumerate(header_row):
        if _header_to_field(cell) == "packing":
            return i
    return None


def _row_spec(
    data_row: tuple[Any, ...],
    idx_map: dict[int, str],
) -> dict[str, str]:
    spec: dict[str, str] = {}
    for col_idx, field in idx_map.items():
        if col_idx >= len(data_row):
            continue
        if field in ("sr_no", "packing", "source_file"):
            continue
        s = _cell_str(data_row[col_idx])
        if s:
            spec[field] = s
    return spec


async def _run(dry_run: bool) -> None:
    wb_path = REPLACEMENT_DIR / STRAINER_REPLACEMENT_FILENAME
    if not wb_path.is_file():
        raise FileNotFoundError(f"Missing strainer workbook: {wb_path}")

    wb = load_workbook(wb_path, read_only=True, data_only=True)
    if STRAINER_SHEET_NAME not in wb.sheetnames:
        wb.close()
        raise ValueError(f"Sheet {STRAINER_SHEET_NAME!r} not in {wb_path.name}")

    ws = wb[STRAINER_SHEET_NAME]
    rows_iter = ws.iter_rows(values_only=True)
    header_row = next(rows_iter, ())
    model_150 = MODEL_BY_KEY[STRAINER_Y150_KEY]
    model_300 = MODEL_BY_KEY[STRAINER_Y300_KEY]
    idx_map_150 = _build_header_map(header_row, model_150)
    idx_map_300 = _build_header_map(header_row, model_300)
    pressure_idx = _pressure_column_index(header_row)
    packing_idx = _packing_column_index(header_row)
    if packing_idx is None:
        wb.close()
        raise ValueError("Packing column not found in strainer workbook header")
    if not idx_map_150 or not idx_map_300:
        wb.close()
        raise ValueError("No mappable headers for strainer sheet")

    pending: list[tuple[str, dict[str, str], str | None]] = []
    for data_row in rows_iter:
        if not data_row or not any(c is not None and str(c).strip() != "" for c in data_row):
            continue
        catalog_key = _strainer_catalog_key(data_row, pressure_idx)
        idx_map = idx_map_300 if catalog_key == STRAINER_Y300_KEY else idx_map_150
        spec = _row_spec(data_row, idx_map)
        if not spec:
            continue
        packing = _cell_str(data_row[packing_idx]) if packing_idx < len(data_row) else None
        pending.append((catalog_key, spec, packing))

    wb.close()
    logger.info("Parsed %d strainer rows with packing from %s", len(pending), wb_path.name)

    if dry_run:
        for catalog_key, spec, packing in pending[:5]:
            logger.info("  %s packing=%r spec=%s", catalog_key, packing, spec)
        if len(pending) > 5:
            logger.info("  ... and %d more", len(pending) - 5)
        return

    await init_db()
    settings = get_settings()
    client_id = settings.ACTIVE_CLIENT
    stats = {"updated": 0, "unchanged": 0, "unmatched": 0, "missing_packing": 0}

    async with async_session_factory() as session:
        for catalog_key, spec, packing in pending:
            if not packing:
                stats["missing_packing"] += 1
                continue
            rid, reason = await match_catalog_row(catalog_key, spec, session)
            if rid is None:
                stats["unmatched"] += 1
                logger.warning("No match for %s (%s): %s", catalog_key, reason, spec)
                continue
            model = MODEL_BY_KEY[catalog_key]
            result = await session.execute(
                select(model).where(
                    model.client_id == client_id,
                    model.row_id == uuid.UUID(rid),
                )
            )
            row = result.scalar_one_or_none()
            if row is None:
                stats["unmatched"] += 1
                continue
            if (row.packing or "").strip() == packing.strip():
                stats["unchanged"] += 1
                continue
            row.packing = packing
            stats["updated"] += 1

        await session.commit()

    logger.info(
        "Strainer packing backfill: updated=%d unchanged=%d unmatched=%d missing_packing=%d",
        stats["updated"],
        stats["unchanged"],
        stats["unmatched"],
        stats["missing_packing"],
    )


def main() -> None:
    logging.basicConfig(level=logging.INFO)
    p = argparse.ArgumentParser(description="Backfill strainer packing column from replacement workbook")
    p.add_argument("--dry-run", action="store_true")
    args = p.parse_args()
    asyncio.run(_run(dry_run=args.dry_run))


if __name__ == "__main__":
    main()
