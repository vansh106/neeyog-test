"""Load ``docs/Final_Products/*.xlsx`` into ``catalog_fp_*`` tables.

Clears prior rows for ``ACTIVE_CLIENT`` on:
  ``catalog_fp_butterfly_all_products`` and every final-product sheet table
  (Mascon / Needle / NRV / Safety / Sampling / Fittings / etc.).

Does **not** touch operator, SOV, limit switch box, brackets, or positioner.

Usage (from ``backend/``):
  python -m db.import_final_products_catalog --dry-run
  python -m db.import_final_products_catalog
"""

from __future__ import annotations

import argparse
import asyncio
import logging
import re
import sys
import uuid
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import delete
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from core.config import get_settings  # noqa: E402
from core.database import async_session_factory, init_db  # noqa: E402
from db.final_product_models import FINAL_PRODUCT_SHEET_MODELS  # noqa: E402
from db.sheet_models import CatalogButterflyValveRow  # noqa: E402

logger = logging.getLogger(__name__)

DOCS_FINAL = Path(__file__).resolve().parents[2] / "docs" / "Final_Products"


def _specs() -> list[tuple[type, str, str, dict[str, Any]]]:
    from db import final_product_models as fpm

    return [
        (CatalogButterflyValveRow, "Butterfly_Valve_All_Combinations.xlsx", "All Products", {}),
        (fpm.CatalogFpMasconManualTcEndRow, "Mascon_Valve_Products (1) - F (1).xlsx", "Manual – TC End", {}),
        (fpm.CatalogFpMasconManualButtWeldRow, "Mascon_Valve_Products (1) - F (1).xlsx", "Manual – Butt Weld", {}),
        (fpm.CatalogFpMasconPneumaticTcEndRow, "Mascon_Valve_Products (1) - F (1).xlsx", "Pneumatic – TC End", {}),
        (fpm.CatalogFpMasconPneumaticButtWeldRow, "Mascon_Valve_Products (1) - F (1).xlsx", "Pneumatic – Butt Weld", {}),
        (fpm.CatalogFpMasconZdvmLTypeRow, "Mascon_Valve_Products (1) - F (1).xlsx", "ZDV-M – L Type", {}),
        (fpm.CatalogFpMasconZdvmJTypeRow, "Mascon_Valve_Products (1) - F (1).xlsx", "ZDV-M – J Type", {}),
        (fpm.CatalogFpMasconZdvpLTypeRow, "Mascon_Valve_Products (1) - F (1).xlsx", "ZDV-P – L Type", {}),
        (fpm.CatalogFpMasconZdvpJTypeRow, "Mascon_Valve_Products (1) - F (1).xlsx", "ZDV-P – J Type", {}),
        (fpm.CatalogFpMasconPrvRow, "Mascon_Valve_Products (1) - F (1).xlsx", "PRV", {}),
        (fpm.CatalogFpMasconAngleScFlangedRow, "Mascon_Valve_Products (1) - F (1).xlsx", "Angle – Screwed & Flanged", {}),
        (fpm.CatalogFpMasconAngleButtWeldRow, "Mascon_Valve_Products (1) - F (1).xlsx", "Angle – Butt Weld", {}),
        (fpm.CatalogFpMasconAngleTcEndRow, "Mascon_Valve_Products (1) - F (1).xlsx", "Angle – TC End", {}),
        (fpm.CatalogFpMasconSpareDiaphragmRow, "Mascon_Valve_Products (1) - F (1).xlsx", "Spare Diaphragm", {}),
        (fpm.CatalogFpNeedleValveRow, "Needle_Valve_Products.xlsx", "Needle Valve", {}),
        (fpm.CatalogFpNrvInlineCheckRow, "Non_Return_Valve_Products.xlsx", "In Line Check Valve", {}),
        (fpm.CatalogFpNrvWaferCheckRow, "Non_Return_Valve_Products.xlsx", "Wafer Check Valve", {}),
        (fpm.CatalogFpNrvNonSlamRow, "Non_Return_Valve_Products.xlsx", "Non Slam Check Valve", {}),
        (fpm.CatalogFpSafetySvBspRow, "Safety_Valve_Products.xlsx", "SV – Screwed (BSP-F)", {}),
        (fpm.CatalogFpSafetySvTcEndRow, "Safety_Valve_Products.xlsx", "SV – TC End", {}),
        (fpm.CatalogFpSafetySvFlanged150Row, "Safety_Valve_Products.xlsx", "SV – Flanged (ASA #150)", {}),
        (fpm.CatalogFpSamplingSvTcEndRow, "Sampling_Valve_Products.xlsx", "SV – TC End", {}),
        (fpm.CatalogFpSamplingSvOdBaseWeldRow, "Sampling_Valve_Products.xlsx", "SV – OD Base Weld End", {}),
        (fpm.CatalogFpSightGlassDoubleWindowRow, "Sight_Glass_Products.xlsx", "Double Window Sight Glass", {}),
        (fpm.CatalogFpSightGlassInlineIcCastedRow, "Sight_Glass_Products.xlsx", "In-Line SG – IC Casted", {}),
        (fpm.CatalogFpSightGlassInlineSolidFlangeRow, "Sight_Glass_Products.xlsx", "In-Line SG – Solid Flange", {}),
        (fpm.CatalogFpStrainerY150Row, "Strainer_Products.xlsx", "Y Strainer – #150", {}),
        (fpm.CatalogFpStrainerY300Row, "Strainer_Products.xlsx", "Y Strainer – #300", {}),
        (fpm.CatalogFpHoseTuderRow, "Hoses_Products.xlsx", "Tuder Hose", {}),
        (fpm.CatalogFpHoseThunderRow, "Hoses_Products.xlsx", "Thunder Hose", {}),
        (fpm.CatalogFpHosePvcNylonNonToxicRow, "Hoses_Products.xlsx", "PVC Nylon Braided – Non-Toxic", {}),
        (fpm.CatalogFpHosePvcNylonFoodGradeRow, "Hoses_Products.xlsx", "PVC Nylon Braided – Food Grade", {}),
        (fpm.CatalogFpHoseRedSiliconRow, "Hoses_Products.xlsx", "Red Silicon Hose", {}),
        (fpm.CatalogFpHosePuRow, "Hoses_Products.xlsx", "PU Hose", {}),
        (fpm.CatalogFpFittingsSmsNutRow, "Fittings_Products.xlsx", "SMS Nut", {}),
        (fpm.CatalogFpFittingsTriCloverEndRow, "Fittings_Products.xlsx", "Tri-Clover End", {}),
        (fpm.CatalogFpFittingsDinNut11851Row, "Fittings_Products.xlsx", "DIN Nut 11851", {}),
        (fpm.CatalogFpFittingsSwivelNutRow, "Fittings_Products.xlsx", "Swivel Nut", {}),
        (fpm.CatalogFpFittingsFlange150Row, "Fittings_Products.xlsx", "Flange #150", {}),
        (fpm.CatalogFpBallValveCasco1PieceMultiEndRow, "Ball_Valve_Products.xlsx", "Casco – 1-Piece (Multi-End)", {}),
        (fpm.CatalogFpBallValveCasco1PieceFlangedRow, "Ball_Valve_Products.xlsx", "Casco – 1-Piece (Flanged)", {}),
        (fpm.CatalogFpBallValveCasco2PieceRow, "Ball_Valve_Products.xlsx", "Casco – 2-Piece", {}),
        (fpm.CatalogFpBallValveCasco3PieceRow, "Ball_Valve_Products.xlsx", "Casco – 3-Piece", {}),
        (fpm.CatalogFpBallValveCasco3PieceExtStemRow, "Ball_Valve_Products.xlsx", "Casco – 3-Piece Ext. Stem", {}),
        (fpm.CatalogFpBallValveCasco3Piece3WayLPortRow, "Ball_Valve_Products.xlsx", "Casco – 3-Piece 3-Way L-Port", {}),
        (fpm.CatalogFpBallValveUnison1PieceMultiEndRow, "Ball_Valve_Products.xlsx", "Unison – 1-Piece (Multi-End)", {}),
        (fpm.CatalogFpBallValveUnison2PieceIsoPadsRow, "Ball_Valve_Products.xlsx", "Unison – 2-Piece ISO Pads", {}),
        (fpm.CatalogFpBallValveUnison3PieceRow, "Ball_Valve_Products.xlsx", "Unison – 3-Piece", {}),
        (fpm.CatalogFpBallValveUnison3Piece3WayLPortRow, "Ball_Valve_Products.xlsx", "Unison – 3-Piece 3-Way L-Port", {}),
        (fpm.CatalogFpFbvBallTypeRow, "Flush_Bottom_Valve_Products - F (1).xlsx", "FBV – Ball Type", {}),
        (fpm.CatalogFpFbvYTypeRow, "Flush_Bottom_Valve_Products - F (1).xlsx", "FBV – Y Type", {}),
    ]


def _header_to_field(h: object) -> str:
    s = str(h or "").strip().lower().replace("–", "-").replace("″", '"')
    s = re.sub(r"[^a-z0-9]+", "_", s)
    s = re.sub(r"_+", "_", s).strip("_")
    return s


def _field_aliases(model: type) -> dict[str, str]:
    """Excel snake_header -> ORM attribute name (first match wins)."""
    cols = {c.key for c in model.__table__.columns} - {
        "row_id",
        "client_id",
        "created_at",
        "updated_at",
    }
    aliases: dict[str, str] = {}
    if "ball_disc" in cols:
        aliases["disc"] = "ball_disc"
    if "ball" in cols:
        aliases["ball"] = "ball"
    if "valve_size" in cols:
        aliases["size"] = "valve_size"
    if "diaphragm" in cols:
        aliases["diaphragm_moc"] = "diaphragm"
    if "actuator_moc" in cols:
        aliases["actuator_type"] = "actuator_moc"
    if "set_pressure" in cols and "set_pressure_range" not in cols:
        aliases["set_pressure"] = "set_pressure"
    return aliases


def _build_header_map(header_row: tuple[Any, ...], model: type) -> dict[int, str]:
    """0-based column index -> ORM field name."""
    cols = {c.key for c in model.__table__.columns} - {
        "row_id",
        "client_id",
        "created_at",
        "updated_at",
    }
    extra = _field_aliases(model)
    idx_map: dict[int, str] = {}
    for i, cell in enumerate(header_row):
        raw = _header_to_field(cell)
        if not raw:
            continue
        name = extra.get(raw, raw)
        if name == "sr" and i + 1 < len(header_row):
            # "sr" + "no" merged cells rare — skip
            pass
        if name in cols:
            idx_map[i] = name
        elif raw == "sr" and "sr_no" in cols:
            idx_map[i] = "sr_no"
    # Sr No often becomes sr_no from "sr_no" header
    for i, cell in enumerate(header_row):
        t = str(cell or "").strip().lower().replace(" ", "_")
        if t in ("sr_no", "sr.no", "sr_no_") and "sr_no" in cols:
            idx_map[i] = "sr_no"
    return idx_map


def _cell_str(v: Any) -> str | None:
    if v is None:
        return None
    if isinstance(v, float):
        if v == int(v):
            return str(int(v))
        return str(v).rstrip("0").rstrip(".") if "." in str(v) else str(v)
    s = str(v).strip()
    return s if s else None


def _cell_float(v: Any) -> float | None:
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(",", "")
    if not s or s in {"-", "–", "—"}:
        return None
    try:
        return float(s)
    except ValueError:
        return None


def _normalize_valve_size(raw: object | None) -> str | None:
    if raw is None:
        return None
    if isinstance(raw, (int, float)):
        if isinstance(raw, float) and raw.is_integer():
            raw = int(raw)
        if isinstance(raw, int):
            return f"DN{raw}"
        s = str(raw).strip()
        return s if s else None
    s = str(raw).strip()
    if not s:
        return None
    if s.isdigit():
        return f"DN{int(s)}"
    return s


async def _clear_final_tables(session: Any, client_id: str) -> None:
    models: list[type] = [CatalogButterflyValveRow]
    models += [m for _, m in FINAL_PRODUCT_SHEET_MODELS]
    seen: set[type] = set()
    for model in models:
        if model in seen:
            continue
        seen.add(model)
        await session.execute(delete(model).where(model.client_id == client_id))


async def _import_all(dry_run: bool) -> None:
    settings = get_settings()
    client_id = settings.ACTIVE_CLIENT
    specs = _specs()

    if dry_run:
        logger.info("Dry run — would clear + import for client_id=%s", client_id)
        for model, fn, sh, _ in specs:
            path = DOCS_FINAL / fn
            logger.info("  %s sheet=%r exists=%s", path.name, sh, path.is_file())
        return

    await init_db()

    async with async_session_factory() as session:
        await _clear_final_tables(session, client_id)
        total = 0

        for model, filename, sheet_name, static_kw in specs:
            path = DOCS_FINAL / filename
            if not path.is_file():
                logger.warning("Skip missing file: %s", path)
                continue
            wb = load_workbook(path, read_only=True, data_only=True)
            if sheet_name not in wb.sheetnames:
                logger.warning("Skip missing sheet %r in %s", sheet_name, path.name)
                wb.close()
                continue
            ws = wb[sheet_name]
            rows_iter = ws.iter_rows(min_row=1, values_only=True)
            header_row = next(rows_iter, ())
            idx_map = _build_header_map(header_row, model)
            if not idx_map:
                logger.warning("No mappable headers for %s %s", filename, sheet_name)
                wb.close()
                continue

            n = 0
            for data_row in rows_iter:
                kwargs: dict[str, Any] = dict(static_kw)
                kwargs["row_id"] = uuid.uuid4()
                kwargs["client_id"] = client_id
                kwargs["source_file"] = filename
                empty = True
                for col_idx, field in idx_map.items():
                    if col_idx >= len(data_row):
                        continue
                    raw = data_row[col_idx]
                    if field == "sr_no":
                        kwargs[field] = _cell_float(raw)
                    elif field == "valve_size" and model is CatalogButterflyValveRow:
                        kwargs[field] = _normalize_valve_size(raw) or _cell_str(raw)
                    else:
                        kwargs[field] = _cell_str(raw)
                    if kwargs.get(field) not in (None, ""):
                        empty = False
                if empty:
                    continue
                session.add(model(**kwargs))
                n += 1
            total += n
            logger.info("Imported %d rows -> %s (%s)", n, model.__tablename__, sheet_name)
            wb.close()

        await session.commit()
        logger.info("Done. Total new rows: %d", total)


def main() -> None:
    logging.basicConfig(level=logging.INFO)
    p = argparse.ArgumentParser()
    p.add_argument("--dry-run", action="store_true")
    args = p.parse_args()
    asyncio.run(_import_all(dry_run=args.dry_run))


if __name__ == "__main__":
    main()
