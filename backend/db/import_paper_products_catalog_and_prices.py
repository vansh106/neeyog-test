"""Import Paper Products from Wholesale Price List workbook.

Parses worksheet ``Paper Products`` — specs go to ``catalog_fp_paper_products``;
list prices per supplier in ``supplier_product_prices``.

Supplier / brand inference (from sheet vendor table + product text):
- Shant Industries / Eco Paper → Eco TNPL cups
- Sahara Industry / Clarro → Clarro cups, wati, tall
- Sai Swaram / SI → SI ripple, double-wall, Eco SI cups
- Dolphin → DL paper cups
- PNS → export / premium PNS cups
- Neeyog / generic PW sections → containers, lids, plates, bio products

Usage (from ``backend/``):
  python -m db.import_paper_products_catalog_and_prices --dry-run
  python -m db.import_paper_products_catalog_and_prices
"""

from __future__ import annotations

import argparse
import asyncio
import logging
import re
import sys
import uuid
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

logger = logging.getLogger(__name__)

WORKBOOK = "Wholesale Price List Date 16.06.2026.xlsx"
WORKSHEET = "Paper Products"
CATALOG_KEY = "fp_paper_products"
DOCS_FINAL = Path(__file__).resolve().parents[2] / "docs" / "Final_Products"

CAT_ECO_CUPS = "Eco Paper Cups"
CAT_CLARRO_TALL = "Clarro Paper Cups Tall"
CAT_CLARRO_CUPS = "Clarro Paper Cups"
CAT_CLARRO_WATI = "Clarro Paper Wati"
CAT_DOLPHIN_CUPS = "Dolphin Paper Cups"
CAT_EXPORT_CUPS = "Export Paper Cups"
CAT_RIPPLE_BROWN = "Ripple Cups Brown"
CAT_RIPPLE_BLACK = "Ripple Cups Black"
CAT_CP_DW = "CP Double Wall Cups"
CAT_SI_DW = "SI Double Wall Cups"
CAT_HIPS_LIDS = "HIPS Lids"
CAT_PAPER_LIDS = "Paper Lids"
CAT_BAGASSE_LIDS = "Bagasse Lids"
CAT_CONTAINER_KRAFT = "Paper Container Kraft"
CAT_SALAD_KRAFT = "PW Salad Kraft Bowl"
CAT_SALAD_WHITE = "PW Salad White Bowl"
CAT_SALAD_SFP = "Paper Salad SFP Bowl"
CAT_CONTAINER_WHITE = "Paper Container White"
CAT_CONTAINER_110 = "Paper Container 110Dia"
CAT_BIO_SUPER = "Biodegradable Super Paper"
CAT_PAPER_PLATES = "Paper Plates"

ALL_CATEGORIES = (
    CAT_ECO_CUPS,
    CAT_CLARRO_TALL,
    CAT_CLARRO_CUPS,
    CAT_CLARRO_WATI,
    CAT_DOLPHIN_CUPS,
    CAT_EXPORT_CUPS,
    CAT_RIPPLE_BROWN,
    CAT_RIPPLE_BLACK,
    CAT_CP_DW,
    CAT_SI_DW,
    CAT_HIPS_LIDS,
    CAT_PAPER_LIDS,
    CAT_BAGASSE_LIDS,
    CAT_CONTAINER_KRAFT,
    CAT_SALAD_KRAFT,
    CAT_SALAD_WHITE,
    CAT_SALAD_SFP,
    CAT_CONTAINER_WHITE,
    CAT_CONTAINER_110,
    CAT_BIO_SUPER,
    CAT_PAPER_PLATES,
)


@dataclass
class ParsedProduct:
    sr_no: int | None
    category: str
    supplier_name: str
    product_name: str
    gsm: str | None
    size_value: str | None
    size_unit: str | None
    diameter_top: str | None
    diameter_bt: str | None
    height: str | None
    dimensions: str | None
    hsn_code: str | None
    rate_per: str | None
    std_pack: str | None
    pack_unit: str | None
    movement: str | None
    primary_godown: str | None
    rate_incl_gst: float | None
    canonical_sku: str
    raw_description: str


def _cell_str(v: Any) -> str | None:
    if v is None:
        return None
    s = str(v).strip()
    return s or None


def _cell_float(v: Any) -> float | None:
    if v is None or v == "" or v == "-":
        return None
    try:
        n = float(v)
        return n if n > 0 else None
    except (TypeError, ValueError):
        return None


def _extract_hsn(text: str) -> str | None:
    m = re.search(r"HSN(?:\s*CODE)?\s*[-:]?\s*(\d{4,8})", text, re.I)
    return m.group(1) if m else None


def _is_section_header(desc: str, sr: Any) -> bool:
    d = desc.strip()
    if not d:
        return False
    if isinstance(sr, (int, float)):
        return False
    dl = d.lower()
    if "page " in dl or "customisation min moq" in dl:
        return True
    if "hsn" in dl:
        return True
    if dl in ("ripple cup brown", "ripple cup black"):
        return True
    if "sr.no" in dl.replace(" ", ""):
        return True
    return False


def _section_category_supplier(desc: str) -> tuple[str, str] | None:
    dl = desc.strip().lower()
    if "eco paper cups" in dl:
        return CAT_ECO_CUPS, "Shant Industries"
    if "clarro paper cups tall" in dl or "clarro paper cups tall" in dl.replace("  ", " "):
        return CAT_CLARRO_TALL, "Sahara Industry"
    if "clarro paper cups" in dl and "leak" in dl:
        return CAT_CLARRO_CUPS, "Sahara Industry"
    if "clarro" in dl and "wati" in dl:
        return CAT_CLARRO_WATI, "Sahara Industry"
    if "dolphin" in dl:
        return CAT_DOLPHIN_CUPS, "Dolphin"
    if "export" in dl and "product" in dl:
        return CAT_EXPORT_CUPS, "PNS"
    if "si ripple cups" in dl:
        return CAT_RIPPLE_BROWN, "Sai Swaram"
    if dl == "ripple cup brown":
        return CAT_RIPPLE_BROWN, "Sai Swaram"
    if dl == "ripple cup black":
        return CAT_RIPPLE_BLACK, "Sai Swaram"
    if "cp double wall" in dl:
        return CAT_CP_DW, "CP"
    if "si double wall" in dl:
        return CAT_SI_DW, "Sai Swaram"
    if "hips lids" in dl:
        return CAT_HIPS_LIDS, "Neeyog Packaging"
    if "paper lids" in dl and "baggase" not in dl:
        return CAT_PAPER_LIDS, "Neeyog Packaging"
    if "baggase lids" in dl:
        return CAT_BAGASSE_LIDS, "Neeyog Packaging"
    if "container(kraft)" in dl.replace(" ", ""):
        return CAT_CONTAINER_KRAFT, "Neeyog Packaging"
    if "salad kraft bowl" in dl:
        return CAT_SALAD_KRAFT, "Neeyog Packaging"
    if "salad white bowl" in dl:
        return CAT_SALAD_WHITE, "Neeyog Packaging"
    if "salad sfp bowl" in dl:
        return CAT_SALAD_SFP, "Neeyog Packaging"
    if "container(white)" in dl.replace(" ", ""):
        return CAT_CONTAINER_WHITE, "Neeyog Packaging"
    if "container 110dia" in dl:
        return CAT_CONTAINER_110, "Neeyog Packaging"
    if "biodegradable super paper" in dl:
        return CAT_BIO_SUPER, "Neeyog Packaging"
    if "paper plate" in dl and "hsn" in dl:
        return CAT_PAPER_PLATES, "Neeyog Packaging"
    return None


def _infer_supplier(desc: str, default: str) -> str:
    dl = desc.lower()
    if "cllaro" in dl or "clarro" in dl:
        return "Sahara Industry"
    if re.search(r"\bsi\b", dl) or "eco si" in dl:
        return "Sai Swaram"
    if re.search(r"\bdl\b", dl) or "dolphin" in dl:
        return "Dolphin"
    if "pns" in dl:
        return "PNS"
    if "eco" in dl and default == "Shant Industries":
        return "Shant Industries"
    if "cp d/w" in dl or re.search(r"\bcp\b", dl):
        return "CP"
    return default


def _extract_volume(desc: str) -> tuple[str | None, str | None]:
    m = re.search(r"(\d+(?:\.\d+)?)\s*ML\b", desc, re.I)
    if m:
        return f"{m.group(1)}ML", "ml"
    m = re.search(r"(\d+(?:\.\d+)?)\s*OZ\b", desc, re.I)
    if m:
        return f"{m.group(1)}OZ", "oz"
    m = re.search(r"(\d+(?:\.\d+)?)\s*LTR\b", desc, re.I)
    if m:
        return f"{m.group(1)}LTR", "ltr"
    m = re.search(r"(\d+)\s*DIA\b", desc, re.I)
    if m:
        return f"{m.group(1)}D", "dia"
    m = re.search(r"\((\d+)D\)", desc, re.I)
    if m:
        return f"{m.group(1)}D", "dia"
    return None, None


def _clean_product_name(desc: str) -> str:
    s = re.sub(r"\s*\*\s*\d+pcs?", "", desc, flags=re.I).strip()
    s = re.sub(r"\s+", " ", s)
    return s


def _canonical_sku(
    category: str,
    product_name: str,
    size_value: str | None,
    gsm: str | None,
    diameter_top: str | None,
) -> str:
    parts = [
        category.lower(),
        (size_value or "").lower(),
        (gsm or "").lower(),
        (diameter_top or "").lower(),
        re.sub(r"[^a-z0-9]+", "-", product_name.lower()).strip("-"),
    ]
    return "|".join(p for p in parts if p)


def _is_product_row(sr: Any, desc: str, rate: float | None) -> bool:
    if not desc.strip():
        return False
    if isinstance(sr, (int, float)):
        return True
    if isinstance(sr, str) and sr.strip():
        dl = desc.lower()
        if "plate" in dl or rate is not None:
            return True
    return False


def _row_text_and_sr(row: tuple[Any, ...]) -> tuple[Any, str] | None:
    """Return (sr, description) for a data row, or None if empty."""
    a = row[0] if len(row) > 0 else None
    b = row[1] if len(row) > 1 else None
    bs = b.strip() if isinstance(b, str) else ""
    if bs:
        return a, bs
    if isinstance(a, str) and a.strip() and not isinstance(a, (int, float)):
        return None, a.strip()
    return None


def _parse_workbook(path: Path) -> list[ParsedProduct]:
    wb = load_workbook(path, read_only=True, data_only=True)
    if WORKSHEET not in wb.sheetnames:
        raise ValueError(f"Worksheet {WORKSHEET!r} not found")
    ws = wb[WORKSHEET]

    current_category = CAT_ECO_CUPS
    current_supplier = "Shant Industries"
    current_hsn: str | None = "48236900"
    out: list[ParsedProduct] = []
    started = False

    for row in ws.iter_rows(values_only=True):
        if len(row) < 2:
            continue

        if not started:
            title_parts = [
                str(row[i]).strip()
                for i in range(min(3, len(row)))
                if row[i] is not None and str(row[i]).strip()
            ]
            title = " ".join(title_parts).lower()
            if "wholesale paper products" in title:
                started = True
            continue

        parsed_row = _row_text_and_sr(row)
        if not parsed_row:
            continue
        sr, desc = parsed_row

        if _is_section_header(desc, sr):
            section = _section_category_supplier(desc)
            if section:
                current_category, current_supplier = section
            hsn = _extract_hsn(desc)
            if hsn:
                current_hsn = hsn
            continue

        rate = _cell_float(row[7] if len(row) > 7 else None)
        if not _is_product_row(sr, desc, rate):
            continue

        supplier = _infer_supplier(desc, current_supplier)
        gsm = _cell_str(row[2] if len(row) > 2 else None)
        dia_top = _cell_str(row[3] if len(row) > 3 else None)
        dia_bt = _cell_str(row[4] if len(row) > 4 else None)
        height = _cell_str(row[5] if len(row) > 5 else None)

        size_value, size_unit = _extract_volume(desc)
        if not size_value and isinstance(sr, str) and "inch" in sr.lower():
            size_value = sr.strip()
            size_unit = "inch"
        if not dia_top and size_unit == "dia" and size_value:
            dia_top = size_value

        dim_parts = [p for p in (dia_top, dia_bt, height) if p]
        dimensions = " x ".join(dim_parts) if dim_parts else None

        sr_no: int | None = None
        if isinstance(sr, (int, float)):
            sr_no = int(sr) if sr == int(sr) else None

        cleaned = _clean_product_name(desc)
        sku = _canonical_sku(current_category, cleaned, size_value, gsm, dia_top)

        out.append(
            ParsedProduct(
                sr_no=sr_no,
                category=current_category,
                supplier_name=supplier,
                product_name=cleaned,
                gsm=gsm,
                size_value=size_value,
                size_unit=size_unit,
                diameter_top=dia_top,
                diameter_bt=dia_bt,
                height=height,
                dimensions=dimensions,
                hsn_code=current_hsn,
                rate_per=_cell_str(row[8] if len(row) > 8 else None),
                std_pack=_cell_str(row[9] if len(row) > 9 else None),
                pack_unit=_cell_str(row[8] if len(row) > 8 else None),
                movement=_cell_str(row[12] if len(row) > 12 else None),
                primary_godown=_cell_str(row[11] if len(row) > 11 else None),
                rate_incl_gst=rate,
                canonical_sku=sku,
                raw_description=desc,
            )
        )

    wb.close()
    return out


async def _get_or_create_supplier(session: Any, client_id: str, name: str) -> Any:
    from sqlalchemy import func, select

    from db.models import Supplier

    result = await session.execute(
        select(Supplier).where(
            Supplier.client_id == client_id,
            func.lower(Supplier.name) == name.strip().lower(),
        )
    )
    row = result.scalar_one_or_none()
    if row:
        return row
    supplier = Supplier(
        client_id=client_id,
        name=name.strip(),
        primary_category_key=CATALOG_KEY,
        is_active=True,
    )
    session.add(supplier)
    await session.flush()
    logger.info("Created supplier %s", name)
    return supplier


async def _run(dry_run: bool) -> None:
    path = DOCS_FINAL / WORKBOOK
    if not path.is_file():
        raise FileNotFoundError(path)

    products = _parse_workbook(path)
    logger.info("Parsed %d product lines from %s", len(products), WORKSHEET)

    by_cat: dict[str, int] = {}
    for p in products:
        by_cat[p.category] = by_cat.get(p.category, 0) + 1
    for cat in ALL_CATEGORIES:
        logger.info("  %s: %d", cat, by_cat.get(cat, 0))

    if dry_run:
        for p in products[:12]:
            logger.info(
                "  [%s / %s] %s vol=%s gsm=%s price=%s",
                p.category,
                p.supplier_name,
                p.product_name,
                p.size_value,
                p.gsm or "",
                p.rate_incl_gst,
            )
        return

    from sqlalchemy import delete, func, select

    from core.config import get_settings
    from core.database import async_session_factory, init_db
    from db.final_product_models import CatalogFpPaperProductsRow
    from db.models import SupplierProductPrice

    settings = get_settings()
    client_id = settings.ACTIVE_CLIENT
    await init_db()
    async with async_session_factory() as session:
        await session.execute(
            delete(CatalogFpPaperProductsRow).where(CatalogFpPaperProductsRow.client_id == client_id)
        )
        await session.execute(
            delete(SupplierProductPrice).where(SupplierProductPrice.catalog_table == CATALOG_KEY)
        )
        await session.flush()

        sku_to_row_id: dict[str, uuid.UUID] = {}
        price_stats = {"created": 0, "updated": 0, "skipped": 0}

        for p in products:
            row_id = sku_to_row_id.get(p.canonical_sku)
            if row_id is None:
                row_id = uuid.uuid4()
                session.add(
                    CatalogFpPaperProductsRow(
                        row_id=row_id,
                        client_id=client_id,
                        sr_no=float(p.sr_no) if p.sr_no is not None else None,
                        variant_type=p.category,
                        product_name=p.product_name,
                        gsm=p.gsm,
                        size_value=p.size_value,
                        size_unit=p.size_unit,
                        diameter_top=p.diameter_top,
                        diameter_bt=p.diameter_bt,
                        height=p.height,
                        dimensions=p.dimensions,
                        hsn_code=p.hsn_code,
                        rate_per=p.rate_per,
                        std_pack=p.std_pack,
                        pack_unit=p.pack_unit,
                        movement=p.movement,
                        primary_godown=p.primary_godown,
                        canonical_sku=p.canonical_sku,
                        source_file=WORKBOOK,
                    )
                )
                sku_to_row_id[p.canonical_sku] = row_id

            if p.rate_incl_gst is None:
                price_stats["skipped"] += 1
                continue

            supplier = await _get_or_create_supplier(session, client_id, p.supplier_name)
            existing = (
                await session.execute(
                    select(SupplierProductPrice).where(
                        SupplierProductPrice.supplier_id == supplier.id,
                        SupplierProductPrice.catalog_table == CATALOG_KEY,
                        SupplierProductPrice.catalog_row_id == row_id,
                    )
                )
            ).scalar_one_or_none()
            if existing:
                existing.list_price_inr = p.rate_incl_gst
                price_stats["updated"] += 1
            else:
                session.add(
                    SupplierProductPrice(
                        supplier_id=supplier.id,
                        catalog_table=CATALOG_KEY,
                        catalog_row_id=row_id,
                        list_price_inr=p.rate_incl_gst,
                    )
                )
                price_stats["created"] += 1

        await session.commit()

        cat_n = (
            await session.execute(
                select(func.count())
                .select_from(CatalogFpPaperProductsRow)
                .where(CatalogFpPaperProductsRow.client_id == client_id)
            )
        ).scalar_one()
        price_n = (
            await session.execute(
                select(func.count())
                .select_from(SupplierProductPrice)
                .where(SupplierProductPrice.catalog_table == CATALOG_KEY)
            )
        ).scalar_one()
        logger.info("Imported catalog rows=%s supplier_prices=%s", cat_n, price_n)
        logger.info("Price stats: %s", price_stats)


def main() -> None:
    logging.basicConfig(level=logging.INFO, format="%(levelname)s:%(name)s:%(message)s")
    parser = argparse.ArgumentParser(description="Import Paper Products catalog and supplier prices")
    parser.add_argument("--dry-run", action="store_true")
    args = parser.parse_args()
    asyncio.run(_run(dry_run=args.dry_run))


if __name__ == "__main__":
    main()
