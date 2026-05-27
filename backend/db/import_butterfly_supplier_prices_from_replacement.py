"""
Import supplier list prices for Butterfly valve from replacement Excel files.

This updates rows in `supplier_product_prices` for:
  - catalog_table = `butterfly_valve`
  - supplier_id = one of PVH-A01 / PVH-S01 / Alfa Laval / Omval

It is meant to be run after you add/refresh the files in:
  docs/Final_Products/replacement/

Notes:
- The Excel sheets include a "Make" column; we ignore it.
- Valve Size formatting in the Excel is sometimes numeric (e.g. 100),
  while our catalog uses `DN{int}`. We normalize accordingly.
"""

from __future__ import annotations

import argparse
import asyncio
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import logging
import openpyxl
from sqlalchemy import String, cast, delete, func, select
from sqlalchemy import Float as SAFloat
import uuid

from core.config import get_settings
from core.database import async_session_factory, init_db
from db.models import Supplier, SupplierProductPrice
from services import pricing_service

CATALOG_TABLE = "butterfly_valve"


@dataclass(frozen=True)
class ReplacementSource:
    xlsx_path: Path
    supplier_name: str


def _clean_str(v: object | None) -> str | None:
    if v is None:
        return None
    s = str(v).strip()
    return s if s else None


def _normalize_valve_size_for_butterfly(v: object | None) -> str | None:
    """
    Database values are like:
      - DN25, DN100
      - 63.5, 101.6 (non-integers stored as-is)

    Replacement sheets typically provide numeric sizes (e.g. 100).
    """
    if v is None:
        return None

    # Numeric (int/float)
    if isinstance(v, (int, float)):
        if isinstance(v, float) and v.is_integer():
            return f"DN{int(v)}"
        if isinstance(v, int):
            return f"DN{v}"
        # Keep non-integers as a trimmed float string
        s = str(v)
        if "." in s:
            s = s.rstrip("0").rstrip(".")
        return s

    # String
    s = str(v).strip()
    if not s:
        return None

    up = s.upper().replace(" ", "")
    if up.startswith("DN"):
        return "DN" + up[2:]

    # "100" -> DN100
    if s.isdigit():
        return f"DN{int(s)}"

    # "63.5" -> "63.5" ; "76.0" -> "DN76" (integer-valued float)
    try:
        fv = float(s.replace(",", ""))
    except ValueError:
        return s

    if fv.is_integer():
        return f"DN{int(fv)}"

    s2 = str(fv)
    if "." in s2:
        s2 = s2.rstrip("0").rstrip(".")
    return s2


def _normalize_end_connection(v: object | None) -> str | None:
    s = _clean_str(v)
    if not s:
        return None
    if s.strip().upper() == "FULL FLANGED":
        return "Flanged"
    return s


def _normalize_construction(variant_type: object | None, construction: object | None) -> str | None:
    v = _clean_str(variant_type)
    c = _clean_str(construction)
    if not c:
        return c

    # Industrial Butterfly Valve Omval sheet uses a shorter construction label.
    if v and v.strip().lower() == "industrial butterfly valve":
        if c.strip().upper() == "CENTRIC DESIGN - WAFER TYPE":
            return "Centric Design - Wafer Type with Mounting Lugs"

    return c


def _read_rows_from_replacement_xlsx(path: Path) -> list[dict[str, Any]]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        sheet_name = wb.sheetnames[0]
        ws = wb[sheet_name]

        rows_iter = ws.iter_rows(min_row=1, values_only=True)
        header = next(rows_iter, ())
        idx = {str(h).strip(): i for i, h in enumerate(header) if h is not None}

        required = [
            "Variant Type",
            "Construction",
            "Valve Size",
            "End Connection",
            "Pressure",
            "Body",
            "Disc",
            "Seat",
            "Price",
        ]
        missing = [k for k in required if k not in idx]
        if missing:
            raise ValueError(f"{path.name}: missing columns: {missing}. Found headers: {list(idx.keys())}")

        def cell(row: tuple, col: str) -> object | None:
            i = idx.get(col)
            if i is None or i >= len(row):
                return None
            return row[i]

        out: list[dict[str, Any]] = []
        for row in rows_iter:
            if not row or all(v is None or (isinstance(v, str) and not v.strip()) for v in row):
                continue

            price_raw = cell(row, "Price")
            # Keep it; import logic will normalize to float.
            price_val = _clean_str(price_raw)
            if not price_val:
                continue

            out.append(
                {
                    # These keys must match CatalogButterflyValveRow columns.
                    "variant_type": _clean_str(cell(row, "Variant Type")),
                    "construction": _normalize_construction(
                        cell(row, "Variant Type"),
                        cell(row, "Construction"),
                    ),
                    "valve_size": _normalize_valve_size_for_butterfly(cell(row, "Valve Size")),
                    "end_connection": _normalize_end_connection(cell(row, "End Connection")),
                    "pressure": _clean_str(cell(row, "Pressure")),
                    "body": _clean_str(cell(row, "Body")),
                    "ball_disc": _clean_str(cell(row, "Disc")),
                    "seat": _clean_str(cell(row, "Seat")),
                    # pricing_service reserves "price" key
                    "price": price_val,
                }
            )
        return out
    finally:
        wb.close()


async def _find_supplier_id(supplier_name: str, db) -> str:
    q = await db.execute(
        select(Supplier).where(
            Supplier.client_id == get_settings().ACTIVE_CLIENT,
            Supplier.name == supplier_name,
        )
    )
    row = q.scalar_one_or_none()
    if not row:
        raise ValueError(f"Supplier not found: {supplier_name}")
    return str(row.id)


async def _delete_supplier_prices_for_catalog(supplier_id: str, catalog_table: str, db) -> int:
    res = await db.execute(
        delete(SupplierProductPrice).where(
            SupplierProductPrice.supplier_id == uuid.UUID(supplier_id),
            SupplierProductPrice.catalog_table == catalog_table,
        )
    )
    return int(res.rowcount or 0)


def _norm_price_value(raw: object | None) -> float | None:
    if raw is None:
        return None
    try:
        s = str(raw).strip().replace(",", "")
    except Exception:
        return None
    if not s:
        return None
    try:
        return float(s)
    except ValueError:
        return None


async def _match_catalog_row_ids(
    catalog_table: str,
    spec_filters: dict[str, Any],
    db,
    *,
    max_matches: int = 50,
) -> list[str]:
    """
    Match Excel spec filters to butterfly catalog rows.

    Unlike `pricing_service.match_catalog_row`, this returns *all* matching
    catalog row ids (duplicates exist in the catalog).
    """
    model = pricing_service.catalog_model_for_table(catalog_table)
    if model is None:
        return []

    client_id = get_settings().ACTIVE_CLIENT
    stmt = select(model.row_id).where(model.client_id == client_id)

    reserved = frozenset({"price", "discount_override", "__skip__"})
    applied = 0
    for col_name, raw_val in (spec_filters or {}).items():
        if col_name in reserved:
            continue
        if raw_val is None:
            continue
        v_raw = str(raw_val).strip()
        if not v_raw:
            continue

        col = getattr(model, col_name, None)
        if col is None:
            continue

        col_t = model.__table__.columns.get(col_name)
        if col_t is None:
            continue

        if isinstance(col_t.type, SAFloat):
            try:
                stmt = stmt.where(col == float(v_raw.replace(",", "")))
            except ValueError:
                continue
        else:
            stmt = stmt.where(func.lower(func.trim(cast(col, String))) == v_raw.lower())
        applied += 1

    if applied == 0:
        return []

    stmt = stmt.limit(max_matches)
    res = await db.execute(stmt)
    return [str(x) for x in res.scalars().all()]


async def run(*, replacements: list[ReplacementSource], dry_run: bool) -> None:
    await init_db()
    settings = get_settings()
    client_id = settings.ACTIVE_CLIENT

    async with async_session_factory() as db:
        for rep in replacements:
            print(f"\n== Importing prices for supplier: {rep.supplier_name} ==")
            supplier_id = await _find_supplier_id(rep.supplier_name, db)
            rows = _read_rows_from_replacement_xlsx(rep.xlsx_path)
            print(f"Parsed rows: {len(rows)} from {rep.xlsx_path.name}")

            total = len(rows)
            matched_rows = 0
            unmatched_rows = 0
            ambiguous_catalog_matches = 0

            # Preview matching without writing (still does catalog lookups).
            for raw in rows:
                price_f = _norm_price_value(raw.get("price"))
                if price_f is None:
                    unmatched_rows += 1
                    continue
                matched = await _match_catalog_row_ids(CATALOG_TABLE, raw, db)
                if matched:
                    matched_rows += 1
                    if len(matched) > 1:
                        ambiguous_catalog_matches += 1
                else:
                    unmatched_rows += 1

            print(
                "Preview:",
                f"total={total}",
                f"matched={matched_rows}",
                f"unmatched={unmatched_rows}",
                f"ambiguous_catalog_matches={ambiguous_catalog_matches}",
            )

            if dry_run:
                continue

            # Replace semantics: clear the whole supplier's butterfly catalog prices,
            # then reinsert from this replacement sheet.
            deleted = await _delete_supplier_prices_for_catalog(supplier_id, CATALOG_TABLE, db)
            print(f"Deleted existing prices: {deleted} rows (supplier={rep.supplier_name})")

            sid = uuid.UUID(supplier_id)
            inserted = 0
            skipped = 0
            seen_catalog_row_ids: set[str] = set()

            for raw in rows:
                price_f = _norm_price_value(raw.get("price"))
                if price_f is None:
                    skipped += 1
                    continue

                matched = await _match_catalog_row_ids(CATALOG_TABLE, raw, db)
                if not matched:
                    skipped += 1
                    continue

                for rid in matched:
                    if rid in seen_catalog_row_ids:
                        continue
                    seen_catalog_row_ids.add(rid)
                    db.add(
                        SupplierProductPrice(
                            supplier_id=sid,
                            catalog_table=CATALOG_TABLE,
                            catalog_row_id=uuid.UUID(rid),
                            list_price_inr=float(price_f),
                            discount_pct_override=None,
                        )
                    )
                    inserted += 1

            await db.commit()
            print(f"Import: inserted={inserted} skipped={skipped}")


def main() -> None:
    parser = argparse.ArgumentParser(description="Import butterfly supplier prices from replacement Excel files")
    parser.add_argument("--dry-run", action="store_true", help="Preview matches without writing")
    args = parser.parse_args()

    logging.getLogger("sqlalchemy.engine").setLevel(logging.WARNING)
    logging.getLogger("sqlalchemy.engine.Engine").setLevel(logging.WARNING)
    logging.getLogger("sqlalchemy.pool").setLevel(logging.WARNING)

    DOCS_FINAL = Path(__file__).resolve().parents[2] / "docs" / "Final_Products" / "replacement"

    # User-provided supplier mapping.
    replacements = [
        ReplacementSource(
            xlsx_path=DOCS_FINAL / "Aluminium Butterfly Valve PVH Make.xlsx",
            supplier_name="PVH-A01",
        ),
        ReplacementSource(
            xlsx_path=DOCS_FINAL / "Hygienic Butterfly Valve Alfa Laval Make.xlsx",
            supplier_name="Alfa Laval",
        ),
        ReplacementSource(
            xlsx_path=DOCS_FINAL / "Hygienic Butterfly Valve PVH Make.xlsx",
            supplier_name="PVH-S01",
        ),
        ReplacementSource(
            xlsx_path=DOCS_FINAL / "Industrial Butterfly Valve Omval Make.xlsx",
            supplier_name="Omval",
        ),
    ]

    missing = [r.xlsx_path for r in replacements if not r.xlsx_path.is_file()]
    if missing:
        raise FileNotFoundError(f"Missing replacement files: {[p.as_posix() for p in missing]}")

    asyncio.run(run(replacements=replacements, dry_run=bool(args.dry_run)))


if __name__ == "__main__":
    main()

